import asyncio
import io
import logging
import os
import re
from pathlib import Path

from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import Application, CallbackQueryHandler, CommandHandler, ContextTypes, MessageHandler, filters

logging.basicConfig(level=logging.INFO)
MAX_BYTES = 15 * 1024 * 1024
HEADER = re.compile(r'^\s*(?:\[(\d+)\]|(\d+)[.)])\s*(.+)$')
CHOICE = re.compile(r'^\s*(?:([A-Ha-h])[).:]|([1-8])[).])\s+(.+)$')

def extract_text(data: bytes, suffix: str) -> str:
    if suffix in {'.txt', '.md', '.csv', '.tsv'}:
        for enc in ('utf-8-sig', 'utf-16', 'cp1251'):
            try:
                return data.decode(enc)
            except UnicodeError:
                continue
        raise ValueError('Не удалось определить кодировку файла')
    if suffix == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                lines.extend(cell.text for cell in row.cells if cell.text.strip())
        return '\n'.join(lines)
    if suffix == '.pdf':
        import pymupdf
        doc = pymupdf.open(stream=data, filetype='pdf')
        try:
            result = '\n'.join(page.get_text(sort=True) for page in doc)
        finally:
            doc.close()
        if not result.strip():
            raise ValueError('Это PDF-скан без текстового слоя. OCR пока не поддерживается.')
        return result
    raise ValueError('Поддерживаются TXT, MD, CSV, TSV, DOCX, PDF (текстовый) и XLSX')


def normalize_excel_header(value) -> str:
    header = str(value or '').strip().casefold()
    for apostrophe in ('’', '‘', 'ʻ', 'ʼ', '`', '´'):
        header = header.replace(apostrophe, "'")
    header = re.sub(r'[_\-/]+', ' ', header)
    return re.sub(r'\s+', ' ', header)


def is_question_header(header: str) -> bool:
    return any(term in header for term in ('вопрос', 'savol', 'савол', 'question'))


def excel_answer_kind(header: str):
    if any(term in header for term in (
        'неправильн', 'неверн', 'ошибочн', 'noto\'g\'ri', 'notogri', 'нотўғри',
        'incorrect', 'wrong answer', 'wrong option', 'distractor',
    )):
        return 'wrong'
    if any(term in header for term in (
        'правильн', 'верн', 'to\'g\'ri', 'togri', 'тўғри', 'тугри', 'correct answer', 'correct',
    )):
        return 'correct'
    if any(term in header for term in (
        'ответ', 'javob', 'жавоб', 'answer', 'вариант', 'variant', 'option', 'choice',
    )):
        return 'answer'
    if re.fullmatch(r'[a-h]', header):
        return 'answer'
    return None


def read_excel_tables(data: bytes):
    """Read question/answer table layouts from all worksheets."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables = []
    try:
        for ws in wb.worksheets:
            values = list(ws.iter_rows(values_only=True))
            header_row = None
            for row_index, row in enumerate(values):
                normalized_row = [normalize_excel_header(value) for value in row]
                question_headers = [
                    i for i, header in enumerate(normalized_row) if is_question_header(header)
                ]
                answer_headers = [
                    i for i, header in enumerate(normalized_row)
                    if i not in question_headers and excel_answer_kind(header)
                ]
                if question_headers and len(answer_headers) >= 2:
                    header_row = row_index
                    break
            if header_row is None:
                continue
            headers = [str(value).strip() if value is not None else '' for value in values[header_row]]
            normalized = [normalize_excel_header(value) for value in headers]
            question_index = next(i for i, header in enumerate(normalized) if is_question_header(header))
            answer_indices = [
                i for i, header in enumerate(normalized)
                if i != question_index and excel_answer_kind(header)
            ]
            if len(answer_indices) < 2:
                continue
            tables.append({
                'sheet': ws.title,
                'headers': headers,
                'question_index': question_index,
                'answer_indices': answer_indices,
                'correct_indices': [i for i in answer_indices if excel_answer_kind(normalized[i]) == 'correct'],
                'rows': values[header_row + 1:],
            })
    finally:
        wb.close()
    return tables


def questions_from_excel_table(table, correct_index: int):
    questions = []
    question_index = table['question_index']
    for row in table['rows']:
        normalized_row = [normalize_excel_header(value) for value in row]
        if (
            any(is_question_header(value) for value in normalized_row)
            and sum(bool(excel_answer_kind(value)) for value in normalized_row) >= 2
        ):
            continue
        def value_at(index):
            if index >= len(row) or row[index] is None:
                return ''
            return str(row[index]).strip()

        question = value_at(question_index)
        answers = [
            ('+' if index == correct_index else '-', value_at(index))
            for index in table['answer_indices']
            if value_at(index)
        ]
        if question and len(answers) >= 2:
            questions.append({'q': question, 'a': answers})
    return questions


def excel_table_correct_index(table):
    correct_indices = table['correct_indices']
    return correct_indices[0] if len(correct_indices) == 1 else None


def excel_table_candidates(table):
    if len(table['correct_indices']) > 1:
        return table['correct_indices']
    return table['answer_indices']

def parse_questions(text: str):
    questions = []
    current = None
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue

        if current is not None:
            choice = CHOICE.match(line)
            if choice:
                current['a'].append(('', choice.group(3).strip()))
                continue

        header = HEADER.match(line)
        if line.startswith('# '):
            header_text = line[2:].strip()
        elif header:
            header_text = header.group(3).strip()
        else:
            header_text = None

        if header_text is not None:
            if current:
                questions.append(current)
            current = {'q': header_text, 'a': []}
            continue

        if current is None:
            continue

        if line.startswith(('+ ', '- ', '* ')):
            mark, answer = line[0], line[2:].strip()
        else:
            choice = CHOICE.match(line)
            if choice:
                mark, answer = '', choice.group(3).strip()
            else:
                mark, answer = '', line
        if answer:
            current['a'].append((mark, answer))

    if current:
        questions.append(current)
    return [q for q in questions if len(q['a']) >= 2]

def convert(questions, mode: str) -> str:
    out = []
    for q in questions:
        marked = [i for i, (mark, _) in enumerate(q['a']) if mark in ('+', '*')]
        if mode == 'marks' and not marked:
            raise ValueError('В некоторых вопросах нет отметки правильного ответа. Выбери режим «Первый вариант» или добавь ключ.')
        out.append('# ' + q['q'])
        for i, (_, answer) in enumerate(q['a']):
            correct = i in marked if mode == 'marks' else i == 0
            out.append(('+ ' if correct else '- ') + answer)
        out.append('')
    return '\n'.join(out).strip() + '\n'

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text('Привет! Я ConvertoBoy 🤖\nПришли файл с тестами (TXT, DOCX, PDF, XLSX, CSV, TSV или MD). Я преобразую вопросы в TXT для «Ассистента».\n\nВажно: в Excel правильный ответ определяется по заголовку столбца; если это неясно, я попрошу выбрать столбец. Для остальных файлов предложу выбрать правило, если правильные ответы не отмечены.')


def excel_selection_keyboard(table):
    candidates = excel_table_candidates(table)
    buttons = [
        InlineKeyboardButton(
            table['headers'][index][:64],
            callback_data=f'excel:select:{candidate_index}',
        )
        for candidate_index, index in enumerate(candidates)
    ]
    return InlineKeyboardMarkup([[button] for button in buttons])


async def deliver_excel(message, questions, filename, status_message):
    result = convert(questions, 'marks')
    output_name = filename + '_Assistant.txt'
    bio = io.BytesIO(result.encode('utf-8-sig'))
    bio.name = output_name
    await message.reply_document(
        document=bio,
        filename=output_name,
        caption=f'Готово: {len(questions)} вопросов. Проверь правильность ответов перед импортом.',
    )
    await status_message.edit_text('Файл готов ✅ Отправь новый документ, чтобы конвертировать ещё один.')


async def document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    doc = update.message.document
    suffix = Path(doc.file_name or '').suffix.lower()
    if doc.file_size and doc.file_size > MAX_BYTES:
        await update.message.reply_text('Максимальный размер файла: 15 МБ.')
        return
    if suffix not in {'.txt','.md','.csv','.tsv','.docx','.pdf','.xlsx'}:
        await update.message.reply_text('Этот формат пока не поддерживается. Пришли TXT, DOCX, PDF, XLSX, CSV, TSV или MD.')
        return
    msg = await update.message.reply_text('Читаю файл…')
    try:
        tgfile = await doc.get_file()
        data = bytes(await tgfile.download_as_bytearray())
        if suffix == '.xlsx':
            tables = await asyncio.to_thread(read_excel_tables, data)
            if not tables:
                raise ValueError(
                    'Не нашёл таблиц с заголовками вопроса и как минимум двух столбцов ответов. '
                    'Используй заголовки «Вопрос», «Правильный ответ», «Неправильный ответ» '
                    'или их аналоги.'
                )
            questions = []
            pending = []
            for table in tables:
                correct_index = excel_table_correct_index(table)
                if correct_index is None:
                    pending.append(table)
                else:
                    questions.extend(questions_from_excel_table(table, correct_index))
            context.user_data['excel_questions'] = questions
            context.user_data['excel_pending'] = pending
            context.user_data['filename'] = Path(doc.file_name or 'tests').stem
            if pending:
                table = pending[0]
                await msg.edit_text(
                    f"На листе «{table['sheet']}» не удалось однозначно определить правильный ответ. "
                    'Выбери столбец правильного ответа:',
                    reply_markup=excel_selection_keyboard(table),
                )
                return
            if not questions:
                raise ValueError('В таблицах не найдено строк с вопросом и как минимум двумя ответами.')
            await deliver_excel(
                update.message,
                questions,
                context.user_data['filename'],
                msg,
            )
            return
        else:
            text = await asyncio.to_thread(extract_text, data, suffix)
            questions = parse_questions(text)
        if not questions:
            raise ValueError('Не нашёл вопросы с вариантами. Для текстовых файлов используй нумерацию или #.')
        context.user_data['questions'] = questions
        context.user_data['filename'] = Path(doc.file_name or 'tests').stem
        await msg.edit_text(f'Нашёл вопросов: {len(questions)}.\nКак определить правильные ответы?', reply_markup=InlineKeyboardMarkup([
            [InlineKeyboardButton('По знаку + или *', callback_data='marks')],
            [InlineKeyboardButton('Первый вариант (предположение)', callback_data='first')]
        ]))
    except Exception as exc:
        logging.exception('Conversion error')
        await msg.edit_text(f'Не удалось обработать файл: {exc}')

async def choose(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    if query.data.startswith('excel:select:'):
        pending = context.user_data.get('excel_pending', [])
        if not pending:
            await query.edit_message_text('Отправь Excel-файл ещё раз: предыдущие данные недоступны.')
            return
        table = pending.pop(0)
        candidates = excel_table_candidates(table)
        try:
            candidate_index = int(query.data.rsplit(':', 1)[1])
            correct_index = candidates[candidate_index]
        except (ValueError, IndexError):
            await query.edit_message_text('Не удалось определить выбранный столбец. Отправь Excel-файл ещё раз.')
            return
        questions = context.user_data.setdefault('excel_questions', [])
        questions.extend(questions_from_excel_table(table, correct_index))
        if pending:
            next_table = pending[0]
            await query.edit_message_text(
                f"На листе «{next_table['sheet']}» не удалось однозначно определить правильный ответ. "
                'Выбери столбец правильного ответа:',
                reply_markup=excel_selection_keyboard(next_table),
            )
            return
        if not questions:
            await query.edit_message_text('В таблицах не найдено строк с вопросом и как минимум двумя ответами.')
            return
        try:
            await deliver_excel(
                query.message,
                questions,
                context.user_data.get('filename', 'tests'),
                query.message,
            )
        except Exception as exc:
            logging.exception('Excel conversion error')
            await query.message.reply_text(f'Не удалось обработать файл: {exc}')
        return

    questions = context.user_data.get('questions')
    if not questions:
        await query.edit_message_text('Отправь файл ещё раз: предыдущие данные недоступны.')
        return
    try:
        result = convert(questions, query.data)
        filename = context.user_data.get('filename', 'tests') + '_Assistant.txt'
        bio = io.BytesIO(result.encode('utf-8-sig'))
        bio.name = filename
        await query.message.reply_document(document=bio, filename=filename, caption=f'Готово: {len(questions)} вопросов. Проверь правильность ответов перед импортом.')
        await query.edit_message_text('Файл готов ✅ Отправь новый документ, чтобы конвертировать ещё один.')
    except Exception as exc:
        await query.message.reply_text(str(exc))

def main():
    token = os.getenv('BOT_TOKEN')
    if not token:
        raise RuntimeError('Укажи BOT_TOKEN в переменных окружения. Не публикуй токен.')
    app = Application.builder().token(token).build()
    app.add_handler(CommandHandler('start', start))
    app.add_handler(MessageHandler(filters.Document.ALL, document))
    app.add_handler(CallbackQueryHandler(choose, pattern='^(marks|first|excel:select:\\d+)$'))
    app.run_polling(drop_pending_updates=True)

if __name__ == '__main__':
    main()
